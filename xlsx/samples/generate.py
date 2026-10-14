"""Regenerates samples/book.xlsx (openpyxl 3.1.5): two sheets with number,
percent, currency and date formats, a merged title, fills, borders, bold
headers, a wide column, a hidden column and a hidden third sheet.

    python3 -m venv /tmp/v && /tmp/v/bin/pip install openpyxl==3.1.5
    /tmp/v/bin/python examples/previewers/xlsx/samples/generate.py
"""

import datetime
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

wb = Workbook()
ws = wb.active
ws.title = "Revenue"
ws.merge_cells("A1:E1")
ws["A1"] = "Quarterly revenue by region"
ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="0A74D9")
ws["A1"].alignment = Alignment(horizontal="center")
head = ["Region", "Q1", "Q2", "Growth", "Updated"]
thin = Side(style="thin", color="999999")
for c, text in enumerate(head, start=1):
    cell = ws.cell(row=2, column=c, value=text)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="E7E6E6")
    cell.border = Border(bottom=Side(style="medium", color="000000"))
rows = [("North", 120000, 140000), ("South", 95000, 130000), ("West", 80000, 88000)]
for r, (region, q1, q2) in enumerate(rows, start=3):
    ws.cell(row=r, column=1, value=region)
    for c, v in ((2, q1), (3, q2)):
        ws.cell(row=r, column=c, value=v).number_format = '"$"#,##0'
    ws.cell(row=r, column=4, value=(q2 - q1) / q1).number_format = "0.0%"
    ws.cell(row=r, column=5, value=datetime.date(2026, 9, 30)).number_format = "yyyy-mm-dd"
    for c in range(1, 6):
        ws.cell(row=r, column=c).border = Border(bottom=thin)
ws.column_dimensions["A"].width = 18
ws.column_dimensions["E"].width = 14
ws["F3"] = "internal note"
ws.column_dimensions["F"].hidden = True

notes = wb.create_sheet("Notes")
notes["A1"] = "Figures are unaudited."
notes["A1"].font = Font(italic=True, color="C00000")
notes["A2"] = 1234567.891
notes["A2"].number_format = "#,##0.00"
wb.create_sheet("Scratch").sheet_state = "hidden"

wb.properties.created = wb.properties.modified = datetime.datetime(2026, 10, 14)
wb.properties.creator = wb.properties.lastModifiedBy = ""
wb.save(pathlib.Path(__file__).with_name("book.xlsx"))
